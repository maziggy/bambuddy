import csv
import io
from datetime import datetime
from typing import Any

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from backend.app.core.printer_scope import PrinterScope
from backend.app.models.archive import PrintArchive
from backend.app.models.print_log import INCLUDE_OPEN_RUNS, PrintLogEntry

# What one print of an archive has of its own (archive redesign): in the
# export, its run's. Everything else is the file's, the archive's.
_PER_PRINT_FIELDS = frozenset(
    {
        "status",
        "quantity",
        "printer_id",
        "printer_name",
        "plate_id",
        "filament_type",
        "filament_used_grams",
        "cost",
        "energy_kwh",
        "energy_cost",
        "wear_cost",
        "failure_reason",
        "user_verdict",
        "duration_seconds",
        "started_at",
        "completed_at",
    }
)


class ExportService:
    """Service for exporting archive data to CSV/Excel formats."""

    # Default fields to export
    DEFAULT_FIELDS = [
        "id",
        "run_id",
        "print_name",
        "filename",
        "status",
        "quantity",
        "printer_id",
        "printer_name",
        "plate_id",
        "project_name",
        "filament_type",
        "filament_used_grams",
        "print_time_seconds",
        "layer_height",
        "nozzle_diameter",
        "bed_temperature",
        "nozzle_temperature",
        "total_layers",
        "cost",
        "designer",
        "tags",
        "notes",
        "failure_reason",
        "user_verdict",
        "duration_seconds",
        "started_at",
        "completed_at",
        "created_at",
    ]

    # Field labels for headers
    FIELD_LABELS = {
        "id": "ID",
        "run_id": "Print ID",
        "printer_name": "Printer",
        "plate_id": "Plate",
        "user_verdict": "Outcome",
        "duration_seconds": "Duration (s)",
        "energy_kwh": "Energy (kWh)",
        "print_name": "Print Name",
        "filename": "Filename",
        "status": "Status",
        "quantity": "Items Printed",
        "printer_id": "Printer ID",
        "project_name": "Project",
        "filament_type": "Filament Type",
        "filament_used_grams": "Filament (g)",
        "print_time_seconds": "Print Time (s)",
        "layer_height": "Layer Height (mm)",
        "nozzle_diameter": "Nozzle (mm)",
        "bed_temperature": "Bed Temp (°C)",
        "nozzle_temperature": "Nozzle Temp (°C)",
        "total_layers": "Total Layers",
        "cost": "Cost",
        "energy_cost": "Energy Cost",
        "wear_cost": "Wear Cost",
        "designer": "Designer",
        "tags": "Tags",
        "notes": "Notes",
        "failure_reason": "Failure Reason",
        "started_at": "Started At",
        "completed_at": "Completed At",
        "created_at": "Created At",
    }

    @classmethod
    def unknown_fields(cls, fields: list[str]) -> list[str]:
        """The requested columns the export does not have: only its own
        columns are exported, never other attributes of an archive (its
        one-tap confirm token is a credential)."""
        return [f for f in fields if f not in cls.FIELD_LABELS]

    def __init__(self, db: AsyncSession):
        self.db = db

    async def export_archives(
        self,
        format: str = "csv",
        fields: list[str] | None = None,
        printer_id: int | None = None,
        project_id: int | None = None,
        status: str | None = None,
        date_from: datetime | None = None,
        date_to: datetime | None = None,
        search: str | None = None,
        visible_to_user_id: int | None = None,
        printer_scope: PrinterScope | None = None,
    ) -> tuple[bytes, str, str]:
        """Export archives to CSV or Excel format, one row per print.

        An archive is a file, each print of it a run (archive redesign): a row
        has the archive's file columns and one run's own (``_PER_PRINT_FIELDS``).
        An archive printed before runs were logged is one row of its own
        fields; one never printed, one row with empty print columns.

        ``printer_id``, ``status`` and the dates pick prints (a run's start,
        an archive without runs by when it was created); ``project_id``,
        ``search`` and ``visible_to_user_id`` pick archives. A caller limited
        to some printers gets only the prints on theirs.

        Returns:
            Tuple of (file_bytes, filename, content_type)
        """
        # Soft-deleted archives (#1343) are excluded: this export is the list
        # the user is looking at, saved to a file, and that list hides them —
        # an export that silently contains rows the UI says are gone is worse
        # than useless for reconciling anything (#2731).
        query = (
            select(PrintArchive)
            .options(selectinload(PrintArchive.project), selectinload(PrintArchive.printer))
            .where(PrintArchive.deleted_at.is_(None))
            .order_by(PrintArchive.created_at.desc())
        )
        if project_id:
            query = query.where(PrintArchive.project_id == project_id)
        if visible_to_user_id is not None:
            query = query.where(PrintArchive.created_by_id == visible_to_user_id)
        from backend.app.core.row_visibility import archive_scope_clause

        if printer_scope is not None and (clause := archive_scope_clause(printer_scope)) is not None:
            query = query.where(clause)
        if search:
            like_pattern = f"%{search}%"
            query = query.where(
                (PrintArchive.print_name.ilike(like_pattern))
                | (PrintArchive.filename.ilike(like_pattern))
                | (PrintArchive.tags.ilike(like_pattern))
                | (PrintArchive.notes.ilike(like_pattern))
                | (PrintArchive.designer.ilike(like_pattern))
            )
        archives = list((await self.db.execute(query)).scalars().all())

        # Every archive's runs, a print running now included; the caller's
        # only, when limited. By the archive query itself, not a list of its
        # ids: a farm's archives outnumber what one statement may bind.
        printed: set[int] = set()
        runs: dict[int, list[PrintLogEntry]] = {}
        if archives:
            archive_ids = query.with_only_columns(PrintArchive.id).order_by(None)
            all_runs = (
                await self.db.scalars(
                    select(PrintLogEntry)
                    .where(PrintLogEntry.archive_id.in_(archive_ids))
                    .order_by(PrintLogEntry.started_at.desc().nulls_last(), PrintLogEntry.id.desc())
                    .execution_options(**{INCLUDE_OPEN_RUNS: True})
                )
            ).all()
            for run in all_runs:
                if run.archive_id is None:
                    continue
                printed.add(run.archive_id)
                if printer_scope is None or printer_scope.allows(run.printer_id):
                    runs.setdefault(run.archive_id, []).append(run)

        statuses = {s.strip() for s in status.split(",") if s.strip()} if status else set()

        def picked(printer: int | None, row_status: str | None, when: datetime | None) -> bool:
            if printer_id and printer != printer_id:
                return False
            if statuses and row_status not in statuses:
                return False
            if date_from and (when is None or when < date_from):
                return False
            return not (date_to and (when is None or when > date_to))

        export_fields = fields if fields else self.DEFAULT_FIELDS
        rows = []
        for archive in archives:
            if archive.id in runs:
                for run in runs[archive.id]:
                    if picked(run.printer_id, run.status, run.started_at or run.created_at):
                        rows.append(self._archive_to_row(archive, export_fields, run))
            elif archive.id not in printed:
                # Printed before runs were logged, or never: the archive's own.
                if picked(archive.printer_id, archive.status, archive.created_at):
                    rows.append(self._archive_to_row(archive, export_fields))
            elif not (printer_id or status or date_from or date_to):
                # Printed only on printers outside the caller's scope: the file.
                rows.append(self._archive_to_row(archive, export_fields, hide_prints=True))

        # Generate headers
        headers = [self.FIELD_LABELS.get(f, f) for f in export_fields]

        # Generate file
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

        if format == "xlsx":
            file_bytes = self._generate_xlsx(headers, rows, export_fields)
            filename = f"archives_export_{timestamp}.xlsx"
            content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        else:
            file_bytes = self._generate_csv(headers, rows)
            filename = f"archives_export_{timestamp}.csv"
            content_type = "text/csv"

        return file_bytes, filename, content_type

    async def export_stats(
        self,
        format: str = "csv",
        days: int = 30,
        printer_id: int | None = None,
        project_id: int | None = None,
        created_by_id: int | None = None,
        printer_scope: PrinterScope | None = None,
    ) -> tuple[bytes, str, str]:
        """Export statistics summary to CSV or Excel format.

        Args:
            format: Export format ('csv' or 'xlsx')
            days: Number of days to include in stats
            printer_id: Filter by printer
            project_id: Filter by project
            created_by_id: Filter by user who created the print (-1 for no user)

        Returns:
            Tuple of (file_bytes, filename, content_type)
        """
        from backend.app.services.failure_analysis import FailureAnalysisService

        # Get failure analysis data (includes stats)
        analysis_service = FailureAnalysisService(self.db)
        analysis = await analysis_service.analyze_failures(
            days=days,
            printer_id=printer_id,
            project_id=project_id,
            created_by_id=created_by_id,
            printer_scope=printer_scope,
        )

        # Build stats rows
        rows = [
            ["Metric", "Value"],
            ["Period (days)", analysis["period_days"]],
            ["Total Prints", analysis["total_prints"]],
            ["Failed Prints", analysis["failed_prints"]],
            ["Failure Rate (%)", analysis["failure_rate"]],
            [""],
            ["Failures by Reason", ""],
        ]

        for reason, count in analysis["failures_by_reason"].items():
            rows.append([reason, count])

        rows.append([""])
        rows.append(["Failures by Filament", ""])

        for filament, count in analysis["failures_by_filament"].items():
            rows.append([filament, count])

        rows.append([""])
        rows.append(["Failures by Printer", ""])

        for printer, count in analysis["failures_by_printer"].items():
            rows.append([printer, count])

        rows.append([""])
        rows.append(["Weekly Trend", ""])
        rows.append(["Week", "Total", "Failed", "Rate (%)"])

        for week in analysis["trend"]:
            rows.append(
                [
                    week["week_start"],
                    week["total_prints"],
                    week["failed_prints"],
                    week["failure_rate"],
                ]
            )

        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")

        if format == "xlsx":
            file_bytes = self._generate_xlsx_simple(rows)
            filename = f"stats_export_{timestamp}.xlsx"
            content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        else:
            file_bytes = self._generate_csv_simple(rows)
            filename = f"stats_export_{timestamp}.csv"
            content_type = "text/csv"

        return file_bytes, filename, content_type

    def _archive_to_row(
        self,
        archive: PrintArchive,
        fields: list[str],
        run: PrintLogEntry | None = None,
        *,
        hide_prints: bool = False,
    ) -> list[Any]:
        """One row: the archive's file columns, and ``run``'s own (without a
        run, the archive's, as it kept its print before runs; empty with
        ``hide_prints``)."""
        row = []
        for field in fields:
            if field == "project_name":
                value = archive.project.name if archive.project else None
            elif field == "run_id":
                value = run.id if run is not None else None
            elif field in _PER_PRINT_FIELDS and hide_prints:
                value = None
            elif field in _PER_PRINT_FIELDS and run is not None:
                value = getattr(run, field, None)
                # Runs logged before they kept these share their archive's.
                if field in ("filament_type", "quantity") and value is None:
                    value = getattr(archive, field)
            elif field == "printer_name":
                value = archive.printer.name if archive.printer is not None else None
            elif field == "duration_seconds":
                started, completed = archive.started_at, archive.completed_at
                value = (
                    int((completed - started).total_seconds())
                    if started and completed and completed > started
                    else None
                )
            else:
                value = getattr(archive, field, None)
            if isinstance(value, datetime):
                value = value.isoformat()
            row.append(value)
        return row

    def _generate_csv(self, headers: list[str], rows: list[list]) -> bytes:
        """Generate CSV file content."""
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(headers)
        writer.writerows(rows)
        return output.getvalue().encode("utf-8")

    def _generate_csv_simple(self, rows: list[list]) -> bytes:
        """Generate CSV file content from simple rows (no separate headers)."""
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerows(rows)
        return output.getvalue().encode("utf-8")

    def _generate_xlsx(self, headers: list[str], rows: list[list], fields: list[str]) -> bytes:
        """Generate Excel file content."""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError:
            raise ImportError("openpyxl is required for Excel export. Install with: pip install openpyxl")

        wb = Workbook()
        ws = wb.active
        ws.title = "Archives"

        # Header style
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center")

        # Write headers
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

        # Write data
        for row_idx, row in enumerate(rows, 2):
            for col_idx, value in enumerate(row, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)

        # Auto-adjust column widths
        for col_idx, _field in enumerate(fields, 1):
            column_letter = get_column_letter(col_idx)
            max_length = len(headers[col_idx - 1])
            for row in rows:
                cell_value = row[col_idx - 1]
                if cell_value is not None:
                    max_length = max(max_length, len(str(cell_value)))
            ws.column_dimensions[column_letter].width = min(max_length + 2, 50)

        # Freeze header row
        ws.freeze_panes = "A2"

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()

    def _generate_xlsx_simple(self, rows: list[list]) -> bytes:
        """Generate Excel file content from simple rows."""
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
        except ImportError:
            raise ImportError("openpyxl is required for Excel export. Install with: pip install openpyxl")

        wb = Workbook()
        ws = wb.active
        ws.title = "Statistics"

        bold_font = Font(bold=True)

        for row_idx, row in enumerate(rows, 1):
            for col_idx, value in enumerate(row, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                # Bold section headers
                if col_idx == 1 and value and isinstance(value, str) and value.endswith(":"):
                    cell.font = bold_font

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()
